from dataclasses import dataclass
from pathlib import Path
import re


@dataclass(frozen=True)
class Document:
    source: str
    text: str


def load_documents(directory: str | Path) -> list[Document]:
    root = Path(directory)
    if not root.is_dir():
        raise FileNotFoundError(f"Document directory does not exist: {root}")
    documents: list[Document] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            text = path.read_text(encoding="utf-8")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        elif suffix == ".docx":
            from docx import Document as DocxDocument
            text = "\n".join(paragraph.text for paragraph in DocxDocument(str(path)).paragraphs)
        else:
            continue
        if text.strip():
            documents.append(Document(str(path), re.sub(r"\s+", " ", text).strip()))
    return documents


def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    words = text.split()
    if not words:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = start
        characters = 0
        while end < len(words) and (end == start or characters + len(words[end]) + 1 <= size):
            characters += len(words[end]) + (1 if end > start else 0)
            end += 1
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        retained = 0
        next_start = end
        while next_start > start and retained < overlap:
            next_start -= 1
            retained += len(words[next_start]) + 1
        start = next_start
    return chunks
